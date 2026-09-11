"""Read the supplied archive without editing it; retain row provenance in local JSON."""
import json
import hashlib
import sys
from pathlib import Path
from collections import Counter
import openpyxl

source = Path(sys.argv[1] if len(sys.argv) > 1 else '/Users/hansen/Downloads/2026.9.10 存档.xlsx')
root = Path(__file__).resolve().parents[1]
workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
brands = []
for row, values in enumerate(workbook['品牌'].values, 1):
    if row == 1 or not values[0]:
        continue
    code, name = [str(v).strip() for v in values[:2]]
    brands.append({'id': code, 'name': name, 'sourceRow': row, 'undefined': code.lower() == 'undefined'})
translations = ['儿童手工', '成人手工', '综合手工', '纸艺', '烘焙', '派对', '礼赠', '儿童趣味小物', '成人趣味小物', '家居日用', '节庆装饰', '家居装饰', '玩具与游戏', '文具', '美妆', '发饰', '厨具', 'Cozy Craftworks']
categories, subcategories = [], []
for row, values in enumerate(workbook['品类'].values, 1):
    if row == 1 or not values[0]:
        continue
    code, name, parent, parent_name = values
    item = {'id': str(code).strip(), 'name': str(name).strip(), 'sourceRow': row}
    if parent is None:
        item['label'] = [translations[int(code) - 1], item['name']]
        categories.append(item)
    else:
        item['parent'] = str(parent).strip()
        subcategories.append(item)
for collection in [brands, categories, subcategories]:
    duplicates = [key for key, n in Counter(i['id'] for i in collection).items() if n > 1]
    if duplicates:
        raise ValueError(f'Duplicate identifiers: {duplicates}')
assert all(s['parent'] in {c['id'] for c in categories} for s in subcategories)
output = {'source': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'brandRange': f'品牌!A2:B{workbook["品牌"].max_row}', 'categoryRange': f'品类!A2:D{workbook["品类"].max_row}',
          'relationshipStatus': 'Brand-to-category and brand ownership are not supplied in this workbook.',
          'brands': brands, 'categories': categories, 'subcategories': subcategories}
(root / 'public/data/catalog.json').write_text(json.dumps(output, ensure_ascii=False, indent=2) + '\n')
print(f'Imported {len(brands)} archive records, {len(categories)} categories, {len(subcategories)} subcategories. Original workbook unchanged.')
